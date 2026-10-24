"""Validate CSV/XLSX inventories, preview, then atomically create new products."""
import csv
import io
import re
from pathlib import Path
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from ..core import Product
from ..models import ActionType
from ..schemas import ImportReport, ImportRowError, ProductCreate, ProductOut
from .common import audit, transaction
from .inventory import add_product

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
COLUMNS = ("sku", "name", "system_quantity", "unit_price", "low_stock_threshold")


def read_rows(filename, content):
    if not content:
        raise HTTPException(422, "The file is empty.")
    if len(content) > MAX_BYTES:
        raise HTTPException(413, "Upload must be at most 5 MiB.")
    extension = Path(filename or "").suffix.lower()
    workbook = None
    try:
        if extension == ".csv":
            rows = csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True)
        elif extension == ".xlsx":
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 25 * 1024 * 1024:
                    raise HTTPException(413, "Expanded workbook exceeds 25 MiB.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
        else:
            raise HTTPException(415, "Use a UTF-8 .csv or .xlsx file. Legacy .xls is not supported.")
        header = next(rows, None)
        if header is None:
            raise HTTPException(422, "The file has no header row.")
        header = [str(cell).strip() if cell is not None else "" for cell in header]
        if len(header) != len(COLUMNS) or set(header) != set(COLUMNS):
            raise HTTPException(422, f"Headers must contain exactly: {', '.join(COLUMNS)}.")
        result = []
        for row_number, row in enumerate(rows, start=2):
            if row_number > MAX_ROWS + 1:
                raise HTTPException(413, f"Import supports at most {MAX_ROWS} data rows.")
            if all(value is None or str(value).strip() == "" for value in row):
                continue
            result.append((row_number, row))
        return header, result
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, "Cannot read the file. Use valid UTF-8 CSV or an unencrypted XLSX workbook.") from error
    finally:
        if workbook is not None:
            workbook.close()


def parse_products(header, rows):
    products, errors, seen = [], [], set()
    for row_number, values in rows:
        try:
            if len(values) != len(header):
                raise ValueError("Row must contain exactly five columns.")
            record = {key: str(value).strip() if value is not None else "" for key, value in zip(header, values)}
            if any(value.startswith("=") for value in record.values()):
                raise ValueError("Formula cells are not supported; provide values instead.")
            for key in ("system_quantity", "low_stock_threshold"):
                if not re.fullmatch(r"[+-]?\d+", record[key]):
                    raise ValueError(f"{key} must be a whole number.")
                record[key] = int(record[key])
            product = ProductCreate.model_validate(record)
            if product.sku in seen:
                raise ValueError(f"Duplicate SKU in file: {product.sku}.")
            seen.add(product.sku)
            products.append((row_number, product))
        except ValidationError as error:
            message = "; ".join(f"{'.'.join(map(str, item['loc']))}: {item['msg']}" for item in error.errors())
            errors.append(ImportRowError(row=row_number, message=message))
        except ValueError as error:
            errors.append(ImportRowError(row=row_number, message=str(error)))
    return products, errors


def import_inventory(db, filename, content, actor, dry_run=True):
    header, rows = read_rows(filename, content)
    candidates, errors = parse_products(header, rows)
    if not rows:
        errors.append(ImportRowError(row=2, message="Add at least one product row."))
    try:
        with transaction(db):
            # Query in batches to stay below SQLite parameter limits.
            skus = [product.sku for _, product in candidates]
            existing = set()
            for start in range(0, len(skus), 500):
                existing.update(db.scalars(select(Product.sku).where(Product.sku.in_(skus[start:start + 500]))))
            for row_number, product in candidates:
                if product.sku in existing:
                    errors.append(ImportRowError(row=row_number, message=f"SKU already exists: {product.sku}. Use stock movements to change stock."))
            report = ImportReport(dry_run=dry_run, valid=not errors, total_rows=len(rows), errors=errors,
                                  preview=[product for _, product in candidates])
            if errors and not dry_run:
                raise HTTPException(422, report.model_dump(mode="json"))
            if not errors and not dry_run:
                for _, product in candidates:
                    created = add_product(db, product, actor)
                    report.products.append(ProductOut.model_validate(created))
                report.created_count = len(candidates)
                audit(db, actor, ActionType.INVENTORY_IMPORTED, None,
                      {"created_count": report.created_count, "skus": skus})
        return report
    except IntegrityError as error:
        raise HTTPException(409, "An SKU was created concurrently or a database constraint failed. Nothing was imported; preview again.") from error
